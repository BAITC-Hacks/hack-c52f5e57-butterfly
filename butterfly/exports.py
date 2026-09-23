"""Local, Unicode-safe protocol exports. Exports never change review state."""

from io import BytesIO
import json
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer


FONT_DIR = Path(__file__).parent / "assets" / "fonts"
for name, filename in [("Butterfly", "DejaVuSans.ttf"), ("ButterflyBold", "DejaVuSans-Bold.ttf")]:
    if name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(name, str(FONT_DIR / filename)))


def _text(value, default="Не указано"):
    return str(value) if value is not None and str(value).strip() else default


def _clock(value):
    if value is None:
        return "—"
    seconds = max(0, int(float(value)))
    return f"{seconds // 60:02}:{seconds % 60:02}"


def _blocks(meeting):
    """A flat stream lets long evidence and transcript paragraphs paginate."""
    yield "brand", "BUTTERFLY / ПРОТОКОЛ ВСТРЕЧИ"
    yield "title", _text(meeting.get("title"), "Встреча")
    yield "meta", f"Дата: {_text(meeting.get('date'))}"
    mode = meeting.get("mode")
    if mode == "demo":
        yield "notice", "ДЕМО · синтетическая встреча"
    elif mode == "manual":
        yield "notice", "Источник: вставленная пользователем стенограмма"
    if meeting.get("status") != "completed":
        yield "notice", "Черновик · требуется проверка"
    else:
        yield "meta", "Проверено пользователем"
    yield "heading", "Краткое содержание"
    yield "body", _text(meeting.get("summary"), "Краткое содержание ещё не подготовлено.")
    yield "heading", "Поручения"
    actions = meeting.get("actions") or []
    if not actions:
        yield "body", "Поручения ещё не добавлены или не извлечены."
    for index, action in enumerate(actions, 1):
        yield "action", f"{index}. {_text(action.get('text'), 'Поручение')}"
        yield "body", f"Ответственный: {_text(action.get('owner'))} · Срок: {_text(action.get('deadline'))}"
        yield "notice" if action.get("status") != "confirmed" else "meta", (
            "НЕ ПОДТВЕРЖДЕНО" if action.get("status") != "confirmed" else "Подтверждено пользователем"
        )
        yield "quote", "Основание: " + _text(action.get("evidence"), "Цитата не указана — проверьте источник.")
    yield "heading", "Стенограмма"
    if not meeting.get("transcript"):
        yield "body", "Стенограмма пока отсутствует. Обработка аудио не подтверждена."
    for segment in meeting.get("transcript") or []:
        start, end = _clock(segment.get("start")), _clock(segment.get("end"))
        yield "meta", f"{start}–{end} · {_text(segment.get('speaker'), 'Говорящий не указан')}"
        yield "body", _text(segment.get("text"), "")


def _pdf(meeting):
    output = BytesIO()
    styles = {}
    for name, size, color, bold in [
        ("brand", 8, "#24715e", True), ("title", 22, "#172d29", True),
        ("heading", 13, "#172d29", True), ("action", 11, "#172d29", True),
        ("body", 10, "#263b35", False), ("quote", 9, "#55635c", False),
        ("meta", 8, "#65746b", False), ("notice", 9, "#966319", True),
    ]:
        styles[name] = ParagraphStyle(name, fontName="ButterflyBold" if bold else "Butterfly",
                                      fontSize=size, leading=size * 1.5, textColor=colors.HexColor(color),
                                      spaceAfter=8, spaceBefore=8 if name=="heading" else 0,
                                      alignment=TA_LEFT, splitLongWords=True)
    story = []
    for kind, text in _blocks(meeting):
        # Split explicit lines into flowables so a single long transcript is splittable.
        for line in text.splitlines() or [""]:
            story.append(Paragraph(escape(line) or "&#160;", styles[kind]))
        if kind == "quote":
            story.append(Spacer(1, 5))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Butterfly", 8)
        canvas.setFillColor(colors.HexColor("#65746b"))
        canvas.drawString(42, 25, "Butterfly · протокол и поручения")
        canvas.drawRightString(A4[0] - 42, 25, str(doc.page))
        canvas.restoreState()

    doc = SimpleDocTemplate(output, pagesize=A4, leftMargin=42, rightMargin=42,
                           topMargin=40, bottomMargin=45,
                           title=_text(meeting.get("title"), "Встреча"), author="Butterfly")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def _docx(meeting):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.7)
    normal = doc.styles["Normal"]
    normal.font.name = "DejaVu Sans"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    for kind, text in _blocks(meeting):
        if kind == "title":
            p = doc.add_heading(text, 0)
        elif kind == "heading":
            p = doc.add_heading(text, 1)
        else:
            p = doc.add_paragraph(text)
        if kind in {"brand", "action", "notice"}:
            for run in p.runs:
                run.bold = True
        if kind in {"meta", "quote"}:
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor.from_string("65746B")
        if kind == "notice":
            for run in p.runs:
                run.font.color.rgb = RGBColor.from_string("966319")
    section.footer.paragraphs[0].text = "Butterfly · протокол и поручения"
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def render_export(meeting: dict, format: str) -> bytes:
    if format == "json":
        return json.dumps(meeting, ensure_ascii=False, indent=2).encode("utf-8")
    if format == "pdf":
        return _pdf(meeting)
    if format == "docx":
        return _docx(meeting)
    raise ValueError("Unsupported export format")
